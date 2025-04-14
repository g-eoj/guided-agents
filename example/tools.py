# Copyright 2025 g-eoj
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.


from guided_agents import Tool


CHUNK_SIZE = 20000


class DOCXReader(Tool):
    name = "docx_reader"
    description = "This is a tool that converts docx files into markdown."
    inputs = {
        "file_path": {
            "type": "string",
            "description": "Path to a docx file.",
        },
    }
    output_type = "string"

    def forward(self, file_path: str):
        import docx
        import re
        from docx import Document
        text = ""

        document = Document(file_path)
        for x in document.iter_inner_content():
            if isinstance(x, docx.text.paragraph.Paragraph):
                if "heading" in x.style.name.lower():
                    text += "#" * int(x.style.name.lower().split()[-1]) + " "
                if "list" in x.style.name.lower():
                    text += "- "
                text += x.text + "\n"
            elif isinstance(x, docx.table.Table):
                for row in x.rows:
                    for cell in row.cells:
                        text += cell.text + ", "
                    text += "\n"
        text = re.sub(r'[\n]{3,}', "\n\n", text)
        return text


class ExcelReader(Tool):
    name = "excel_reader"
    description = (
        "This is a tool that converts excel files into text. "
        "Each sheet will be followed by two tables of CSVs. "
        "The first table contains the values of the cells in the sheet. "
        "The second table contains the style ID of the cell."
    )
    inputs = {
        "file_path": {
            "type": "string",
            "description": "Path to the excel file.",
        },
    }
    output_type = "string"

    def forward(self, file_path: str):
        import openpyxl

        work_book = openpyxl.open(file_path)
        text = ""
        for sheet in work_book:
            values = ""
            style_ids = ""
            text += str(sheet) + "\n"
            for row in sheet.iter_rows():
                for cell in row:
                    values += str(cell.value) + ", "
                    style_ids += str(cell.style_id) + ", "
                values = values[:-2] + "\n"
                style_ids = style_ids[:-2] + "\n"
            text += values
            text += "\n"
            text += style_ids
            text += "\n"
        text +="\n"
        return text


class FinalAnswerReviewer(Tool):
    name = "final_answer"
    description = "When you are done with your task, use this tool to provide your answer."
    inputs = {
        "answer": {
            "type": "string",
            "description": "The answer to your task."
        },
    }
    output_type = "string"

    def __init__(self, model, question: str, requirements: str):
        super().__init__(self)
        self.model = model
        self.question = question
        self.requirements = requirements

    def forward(self, answer: str):
        messages = [
            {"role": "user", "content":
                [
                    {
                        "type": "text",
                        "text": (
                            f"Check if this answer '{answer}' meets these requirements:\n{self.requirements}"
                            f"\n\nIf the answer '{answer}' does meet the requirements, say the answer. "
                            "If the answer doesn't meet the requirements, say the fixed answer."
                        )
                    }
                ]
            }
        ]
        text = str(self.model(messages=messages).content)
        print(text)
        return text or answer


class GoogleScholarSearchTool(Tool):
    name = "paper_search"
    description = "Search the web for research papers."
    inputs = {
        "query": {"type": "string", "description": "The search query to perform."},
    }
    output_type = "string"

    def __init__(self, api_key):
        super().__init__(self)
        import requests_cache
        self.api_key = api_key
        self.headers = {'X-API-KEY': api_key}
        self.url = "https://google.serper.dev/scholar"
        self.session = requests_cache.CachedSession("google_search_cache", allowable_methods=["POST"])

    def forward(self, query: str) -> str:
        result = self.session.post(url=self.url, data={"q": query}, headers=self.headers)
        result.raise_for_status()
        result = result.json()["organic"]
        text = "\n"
        for r in result:
            pdf_url = r.get("pdfUrl", None)
            html_url = r.get("htmlUrl", None)
            link_url = r.get("link", None)
            link = pdf_url or html_url or link_url
            text += f"- [{r['title']}]({link}): {r.get('snippet', '')}...\n\tPublication Information: {r.get('publicationInfo', '')}\n"
        return text


class GoogleSearchTool(Tool):
    name = "web_search"
    description = "Search the web for links."
    inputs = {
        "query": {"type": "string", "description": "The search query to perform."},
    }
    output_type = "string"

    def __init__(self, api_key):
        super().__init__(self)
        import requests_cache
        self.api_key = api_key
        self.headers = {'X-API-KEY': api_key}
        self.url = "https://google.serper.dev/search"
        self.session = requests_cache.CachedSession("google_search_cache", allowable_methods=["POST"])

    def forward(self, query: str) -> str:
        result = self.session.post(url=self.url, data={"q": query, "num": 8}, headers=self.headers)
        result.raise_for_status()
        result = result.json()["organic"]
        text = "\n"
        for r in result:
            text += f"- [{r['title']}]({r['link']}): {r.get('snippet', '')}\n"
        return text


class MLXAudioTranscribe(Tool):
    name = "audio_transcribe"
    description = "Transcribe speech from audio files."
    inputs = {
        "file_path": {"type": "string", "description": "Audio file to transcribe."},
    }
    output_type = "string"

    def forward(self, file_path: str) -> str:
        import mlx_whisper
        text = mlx_whisper.transcribe(file_path)["text"]
        return text


class PPTXReader(Tool):
    name = "pptx_reader"
    description = "This tool converts pptx files into a list. Each element of the list is text from one slide."
    inputs = {
        "file_path": {"type": "string", "description": "Path to a pptx file."},
    }
    output_type = "array"

    def forward(self, file_path: str) -> str:
        from pptx import Presentation
        presentation = Presentation(file_path)
        slides_text = []
        for slide in presentation.slides:
            text = ""
            element = slide.element
            for x in element.iterdescendants():
                if x.text and hasattr(x, "xml") and x.xml.startswith("<a:r"):
                    text += x.text + "\n"
            slides_text.append(text)
        return slides_text


class TXTReader(Tool):
    name = "txt_reader"
    description = "This is a tool that converts txt files into text."
    inputs = {
        "file_path": {
            "type": "string",
            "description": "Path to the txt file.",
        },
    }
    output_type = "string"

    def forward(self, file_path: str):
        with open(file_path, "r") as h:
            text = "".join(h.readlines())
        return text


class WebReader(Tool):
    name = "web_reader"
    description = "Converts a web page or PDF to markdown then answers your query about the markdown content."
    inputs = {
        "path": {
            "type": "string",
            "description": "The URL or file path to read from."
        },
        "query": {
            "type": "string",
            "description": "Ask a simple but detailed question about what you want to learn."
        }
    }
    output_type = "string"

    def __init__(self, model):
        super().__init__(self)
        self.model = model

    def get_page(self, path:str):
        import re
        import time
        from urllib.parse import urlparse

        import markdownify
        from playwright.sync_api import sync_playwright

        hostname = urlparse(path).hostname
        if hostname and "researchgate.net" in hostname:
            raise ValueError(f"{hostname} is forbidden. Try a different website.")
        md = ""
        with sync_playwright() as playwright:
            browser = playwright.firefox.launch(
                firefox_user_prefs={
                    "pdfjs.disabled": False,
                    "browser.download.open_pdf_attachments_inline": True,
                    "browser.link.open_newwindow": 1,
                }
            )
            page = browser.new_page()
            page.goto(path, wait_until="commit")
            time.sleep(3)
            for frame in page.frames:
                try:
                    # try loading the pdf viewer
                    content = frame.inner_html("id=viewer", timeout=1000)
                except Exception:
                    content = frame.page.inner_html("body")
                md += markdownify.markdownify(
                    content,
                    strip=["a"],
                    heading_style="ATX",
                    table_infer_header=True,
                ) + "\n\n"
            browser.close()
        md = re.sub(r"\n{3,}", "\n\n", md)
        return md

    def forward(self, path: str, query: str):
        try:
            md = self.get_page(path)
        except ValueError as e:
            print(e)
            return str(e)
        notes = []
        chunk_size = CHUNK_SIZE
        chunk_overlap = CHUNK_SIZE * 2
        stage = 0
        for i in range(min(len(md) // chunk_size + 1, 60)):
            windows = []
            chunk = "<chunk_start>" + md[i*chunk_size:(i+1)*chunk_size+chunk_overlap] + "<chunk_end>"
            messages = [
                {"role": "user", "content":
                    [
                        {
                            "type": "text",
                            "text": (
                                f"Look for information relevant to '{query}', using this chunk of markdown:\n\n{chunk}\n\n"
                                f"Now, only using the chunk of markdown, make a note about '{query}' with a short sentence or a list, unless the chunk isn't relevant. "
                                f"If the chunk isn't relevant to '{query}', then say <not_relevant>."
                            )
                        }
                    ]
                }
            ]
            note = self.model(messages=messages).content
            if "<not_relevant>" not in note or notes:
                notes.append(note)
                if len(notes) > 2 or "<not_relevant>" in note:
                    break
        if len(notes) > 1:
            messages = [
                {"role": "user", "content":
                    [
                        {
                            "type": "text",
                            "text": (
                                "Combine these notes. "
                                f"Discard information that is not relevant to '{query}'. "
                                "Check if the notes have contradictory information. "
                                f"If they do, discard the contradictory note with the least information:\n\n{notes}"
                            )
                        }
                    ]
                }
            ]
            answer = self.model(messages=messages).content
            return answer
        elif notes:
            return notes[0]
        else:
            return "No answers found."


class WikipediaReader(Tool):
    name = "wikipedia_reader"
    description = "Reads a Wikipedia article then makes notes on question about the article."
    inputs = {
        "topic": {
            "type": "string",
            "description": "The topic of the Wikipedia article."
        },
        "query": {
            "type": "string",
            "description": "Ask a simple but very specific question about what you want to learn."
        }
    }
    output_type = "string"

    def __init__(self, model):
        super().__init__(self)
        self.model = model

    def forward(self, topic: str, query: str):
        import markdownify
        import requests
        from bs4 import BeautifulSoup, SoupStrainer

        try:
            search_result = requests.get(f"https://api.wikimedia.org/core/v1/wikipedia/en/search/page?q={topic}&limit=1")
            key = search_result.json()["pages"][0]["key"]
            page_html = requests.get(f"https://api.wikimedia.org/core/v1/wikipedia/en/page/{key}/html")
            soup = BeautifulSoup(
                page_html.text, 'html.parser', parse_only=SoupStrainer("body")
            )
            md = markdownify.markdownify(
                str(soup),
                strip=["a"],
                heading_style="ATX",
                table_infer_header=True,
            )
            md = md.split("## See also\n")[0]
            md = md.split("## References\n")[0]

            answers = []
            chunk_size = CHUNK_SIZE
            chunk_overlap = CHUNK_SIZE // 2
            step = 0
            for i in range(len(md) // chunk_size + 1):
                chunk = "<chunk_start>" + md[i*chunk_size:(i+1)*chunk_size+chunk_overlap] + "<chunk_end>"
                messages = [
                    {"role": "user", "content":
                        [
                            {
                                "type": "text",
                                "text": (
                                    f"Look for information relevant to '{query}', using this chunk of markdown:\n\n{chunk}\n\n"
                                    f"Now make a short note of the information in one sentence or list, unless the chunk isn't relevant. "
                                    f"If the chunk isn't relevant to '{query}', then just say <not_relevant>."
                                )
                            }
                        ]
                    }
                ]
                answer = self.model(messages=messages).content
                if "<not_relevant>" not in answer:
                    answers.append(answer)
                    step += 1
                elif step:
                    step += 1
                if step > 2:
                    break
            if answers:
                return f"Notes: {str(answers)}"
            else:
                return "No answers found."
        except Exception as e:
            print(e)
            return(f"No articles found about {topic}.")
